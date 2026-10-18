#!/usr/bin/env python3
"""Export complet Home Assistant -> Excel pour CITADEL CONTROL.
Usage: HA_URL=http://10.0.50.7:8123 HA_TOKEN='...' python3 export_citadel_entities.py
Requires: python3 -m pip install openpyxl websocket-client
"""
import os, json, urllib.request, datetime, collections, ssl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

BASE=os.environ.get('HA_URL','http://10.0.50.7:8123').rstrip('/')
TOKEN=os.environ.get('HA_TOKEN','')
if not TOKEN: raise SystemExit('HA_TOKEN absent. Charger le .env du backend sans afficher son contenu.')

def rest(path):
    req=urllib.request.Request(BASE+path,headers={'Authorization':'Bearer '+TOKEN,'Content-Type':'application/json'})
    with urllib.request.urlopen(req,timeout=45) as response: return json.load(response)

def registries():
    try:
        import websocket
        ws_url=BASE.replace('https://','wss://').replace('http://','ws://')+'/api/websocket'
        ws=websocket.create_connection(ws_url,timeout=25)
        hello=json.loads(ws.recv())
        if hello.get('type')!='auth_required': raise RuntimeError('Handshake WebSocket inattendu')
        ws.send(json.dumps({'type':'auth','access_token':TOKEN}))
        auth=json.loads(ws.recv())
        if auth.get('type')!='auth_ok': raise RuntimeError('Authentification WebSocket refusée')
        output={}
        for idx,(key,command) in enumerate([('areas','config/area_registry/list'),('devices','config/device_registry/list'),('entities','config/entity_registry/list')],start=1):
            ws.send(json.dumps({'id':idx,'type':command}))
            while True:
                msg=json.loads(ws.recv())
                if msg.get('id')==idx:
                    if not msg.get('success'): raise RuntimeError(f'{command}: {msg.get("error")}')
                    output[key]=msg.get('result',[])
                    break
        ws.close()
        return output
    except Exception as exc:
        print('Registres indisponibles, export des états maintenu :',str(exc))
        return {'areas':[],'devices':[],'entities':[]}

def clean(v,limit=900):
    if v is None: return ''
    if isinstance(v,(dict,list,tuple)): v=json.dumps(v,ensure_ascii=False,default=str)
    v=str(v)
    if v.startswith(('=','+','-','@')): v="'"+v
    return v[:limit]

def sheet(wb,name,headers,rows,widths):
    ws=wb.create_sheet(name)
    ws.append(headers)
    for row in rows: ws.append([clean(x,32700) if isinstance(x,(dict,list)) else clean(x,900) for x in row])
    ws.freeze_panes='C2' if len(headers)>4 else 'A2'
    ws.sheet_view.showGridLines=False
    for c,h in enumerate(headers,1):
        cell=ws.cell(1,c);cell.fill=PatternFill('solid',fgColor='14243B');cell.font=Font(color='FFFFFF',bold=True,size=11);cell.alignment=Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(c)].width=widths[c-1] if c-1<len(widths) else 20
    ws.row_dimensions[1].height=29
    if rows:
        tab=Table(displayName='T'+''.join(x for x in name if x.isalnum()),ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
        tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(tab)
    return ws

states=rest('/api/states')
regs=registries()
areas={a.get('area_id'):a.get('name','') for a in regs['areas']}
devices={d.get('id'):d for d in regs['devices']}
ereg={e.get('entity_id'):e for e in regs['entities']}
count=collections.Counter()
rows=[]
for e in sorted(states,key=lambda x:x['entity_id']):
    eid=e['entity_id'];domain=eid.split('.',1)[0];count[domain]+=1
    at=e.get('attributes') or {};r=ereg.get(eid,{})
    device=devices.get(r.get('device_id'),{})
    area_id=r.get('area_id') or device.get('area_id')
    area=areas.get(area_id,'')
    rows.append([eid,domain,at.get('friendly_name') or r.get('name') or '',e.get('state',''),area,area_id or '',r.get('device_id') or '',device.get('name_by_user') or device.get('name') or '',r.get('platform',''),r.get('disabled_by') or '',r.get('hidden_by') or '',at.get('device_class',''),at.get('state_class',''),at.get('unit_of_measurement',''),at.get('supported_features',''),at.get('icon',''),e.get('last_changed',''),e.get('last_updated',''),at])

wb=Workbook();intro=wb.active;intro.title='LIRE AVANT'
notes=[['CITADEL CONTROL — INVENTAIRE HOME ASSISTANT'],['Export UTC',datetime.datetime.now(datetime.timezone.utc).isoformat()],['URL Home Assistant',BASE],['Entités actives retournées par /api/states',len(states)],['Entrées registre des entités',len(regs['entities'])],['Appareils',len(regs['devices'])],['Pièces',len(regs['areas'])],[''],['UTILISATION'],['1. Onglet Entites : liste complète des entités et attributs bruts.'],['2. Onglet Domaines : volumes par type pour prioriser les pages.'],['3. Onglets Pieces, Appareils et Registre : rattachements techniques.'],['4. Ne pas interpréter unknown/unavailable comme des équipements opérationnels.'],['5. Ne pas inventer de valeurs ni de correspondances de pièces.'],['6. Ne pas inclure HA_TOKEN ni de données personnelles dans le code frontend.'],['7. L’inventaire contient potentiellement des identifiants, positions, présences et états privés : ne pas le publier.']]
for row in notes: intro.append(row)
intro.column_dimensions['A'].width=95;intro.column_dimensions['B'].width=48
intro['A1'].font=Font(bold=True,color='FFFFFF',size=16);intro['A1'].fill=PatternFill('solid',fgColor='14243B');intro.row_dimensions[1].height=36
intro['A9'].font=Font(bold=True,color='2860A0',size=13)
headers=['entity_id','domaine','nom','etat','piece','area_id','device_id','appareil','integration','desactive_par','masque_par','device_class','state_class','unite','supported_features','icone','last_changed','last_updated','attributs_json']
sheet(wb,'Entites',headers,rows,[48,18,42,25,24,25,38,35,24,20,20,22,20,16,20,24,28,28,70])
sheet(wb,'Domaines',['domaine','nombre'],sorted(count.items(),key=lambda x:-x[1]),[24,18])
sheet(wb,'Pieces',['area_id','nom','etage','aliases'],[[a.get('area_id'),a.get('name'),a.get('floor_id'),a.get('aliases')] for a in regs['areas']],[35,32,25,60])
sheet(wb,'Appareils',['device_id','nom','piece','fabricant','modele','integration','via_device_id','identifiants'],[[d.get('id'),d.get('name_by_user') or d.get('name'),areas.get(d.get('area_id'),''),d.get('manufacturer'),d.get('model'),d.get('config_entries'),d.get('via_device_id'),d.get('identifiers')] for d in regs['devices']],[38,42,26,28,34,50,38,65])
sheet(wb,'Registre',['entity_id','device_id','area_id','platform','disabled_by','hidden_by','original_name','entity_category'],[[r.get('entity_id'),r.get('device_id'),r.get('area_id'),r.get('platform'),r.get('disabled_by'),r.get('hidden_by'),r.get('original_name'),r.get('entity_category')] for r in regs['entities']],[50,38,28,25,20,20,42,22])
path=os.environ.get('OUTPUT_XLSX','/opt/citadel-control/CITADEL_CONTROL_Entites_Home_Assistant.xlsx')
wb.save(path)
print(f'Export terminé : {path} | {len(states)} entités | {len(regs["areas"])} pièces | {len(regs["devices"])} appareils')
